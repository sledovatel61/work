# -*- coding: utf-8 -*-
"""
UNIT GROUP — финальная презентация (агент 3, v2 после ревью по скринам).
5 слайдов 2400x1350 (16:9) -> PNG -> PDF (13.33"x7.5") + PPTX.

v2 — правки по фидбеку:
  * сетка не «выколи глаз»: убраны сплошные линии, остались только
    угловые маркеры и редкие крестовины-реперы;
  * изображения показываются ЦЕЛИКОМ (contain-fit + размытая подложка,
    без обрезки), яркость поднята;
  * фоновые картинки есть на всех слайдах (титул, структура, компетенции,
    проекты, контакты) и реально видны;
  * в карточках «Геотехника» и «Мичуринский» добавлены сгенерированные фото.
Дизайн-система «Dark Engineering» (агент 2): #0A0E14 / #FF6B00 / #22D3EE,
шрифты Unbounded / Inter / IBM Plex Mono. Тело текста 21–30 px на 2400 px.
"""
import os
from PIL import Image, ImageDraw, ImageFont, ImageFilter, ImageEnhance

BASE = os.path.dirname(os.path.abspath(__file__))
IMG, FNT = os.path.join(BASE, "img"), os.path.join(BASE, "fonts")
W, H = 2400, 1350

BG      = (10, 14, 20)
BG2     = (14, 27, 46)
CARD    = (22, 38, 60)
CARD2   = (17, 30, 48)
LINE    = (30, 52, 80)
ORANGE  = (255, 107, 0)
ORANGE2 = (255, 140, 50)
CYAN    = (34, 211, 238)
GOLD    = (212, 175, 55)
WHITE   = (232, 237, 242)
GREY    = (138, 151, 168)
GREY2   = (110, 122, 138)
FAINT   = (86, 98, 114)

_font_cache = {}
def F(role, size):
    key = (role, size)
    if key in _font_cache: return _font_cache[key]
    if role.startswith("disp"):
        w = {"disp900":900,"disp800":800,"disp700":700,"disp600":600,"disp500":500}[role]
        f = ImageFont.truetype(os.path.join(FNT, "Unbounded[wght].ttf"), size)
        f.set_variation_by_axes([w])
    elif role.startswith("body"):
        w = {"body400":400,"body500":500,"body600":600,"body700":700}[role]
        f = ImageFont.truetype(os.path.join(FNT, "Inter[opsz,wght].ttf"), size)
        f.set_variation_by_axes([28, w])
    else:
        fn = {"mono300":"IBMPlexMono-Regular.ttf","mono400":"IBMPlexMono-Regular.ttf",
              "mono500":"IBMPlexMono-Medium.ttf","mono600":"IBMPlexMono-SemiBold.ttf",
              "mono700":"IBMPlexMono-Bold.ttf"}[role]
        f = ImageFont.truetype(os.path.join(FNT, fn), size)
    _font_cache[key] = f
    return f

# ---------------------------------------------------------------- helpers
def canvas():
    im = Image.new("RGB", (W, H), BG)
    return im, ImageDraw.Draw(im)

def cover_of(pic, bw, bh):
    pr, br = pic.width / pic.height, bw / bh
    if pr > br:                      # картинка шире рамки — режем бока
        nw, nh = int(bh * pr), bh
        off = (nw - bw) // 2
        return pic.resize((nw, nh), Image.LANCZOS).crop((off, 0, off + bw, bh))
    nw, nh = bw, int(bw / pr)        # картинка выше рамки — режем верх/низ
    off = (nh - bh) // 2
    return pic.resize((nw, nh), Image.LANCZOS).crop((0, off, bw, off + bh))

def load_img(name, brightness=1.0):
    p = Image.open(os.path.join(IMG, name)).convert("RGB")
    return ImageEnhance.Brightness(p).enhance(brightness) if brightness != 1.0 else p

def cover_img(im, name, box, brightness=1.0, extra_dark=0):
    x0, y0, x1, y1 = map(int, box)
    im.paste(cover_of(load_img(name, brightness), x1 - x0, y1 - y0), (x0, y0))
    if extra_dark:
        overlay(im, box, (5, 8, 14), extra_dark)

def fit_contain(im, name, box, brightness=1.0, scrim=110):
    """Картинка ЦЕЛИКОМ в box (без кропа): сзади — размытая она же как подложка."""
    x0, y0, x1, y1 = map(int, box)
    bw, bh = x1 - x0, y1 - y0
    pic = load_img(name, brightness)
    fill = cover_of(load_img(name, brightness * 0.75), bw, bh).filter(ImageFilter.GaussianBlur(30))
    fill = Image.blend(fill, Image.new("RGB", (bw, bh), BG), 0.45)
    im.paste(fill, (x0, y0))
    r = min(bw / pic.width, bh / pic.height)
    nw, nh = max(1, int(pic.width * r)), max(1, int(pic.height * r))
    im.paste(pic.resize((nw, nh), Image.LANCZOS), (x0 + (bw - nw) // 2, y0 + (bh - nh) // 2))

def overlay(im, box, color, alpha):
    x0, y0, x1, y1 = map(int, box)
    ov = Image.new("RGBA", (x1 - x0, y1 - y0), color + (alpha,))
    im.paste(Image.alpha_composite(im.crop((x0, y0, x1, y1)).convert("RGBA"), ov).convert("RGB"), (x0, y0))

def vgrad(im, box, c1, c2, alpha=255, flip=False):
    """Полупрозрачный вертикальный градиент поверх картинки (alpha -> 0).
    flip=True — темнее к низу."""
    x0, y0, x1, y1 = map(int, box)
    w, hh = x1 - x0, max(1, y1 - y0)
    band = Image.new("RGBA", (w, hh))
    bd = ImageDraw.Draw(band)
    for y in range(hh):
        t = y / max(1, hh - 1)
        if flip: t = 1 - t
        col = tuple(int(c1[i] + (c2[i] - c1[i]) * t) for i in range(3)) + (int(alpha * (1 - t)),)
        bd.line([(0, y), (w, y)], fill=col)
    im.paste(Image.alpha_composite(im.crop((x0, y0, x1, y1)).convert("RGBA"), band).convert("RGB"), (x0, y0))

def hfade(im, box, color, a0, a1):
    """Горизонтальный градиент затемнения: слева a0 -> справа a1."""
    x0, y0, x1, y1 = map(int, box)
    w = max(2, x1 - x0)
    m = Image.new("L", (w, 1))
    for x in range(w):
        m.putpixel((x, 0), int(a0 + (a1 - a0) * x / (w - 1)))
    m = m.resize((w, y1 - y0))
    im.paste(Image.composite(Image.new("RGB", m.size, color), im.crop((x0, y0, x1, y1)), m), (x0, y0))

def deco(d, plus=True):
    """Деликатный инженерный декор: угловые маркеры + редкие крестовины."""
    for cx, cy, dx, dy in [(42, 42, 1, 1), (W - 42, 42, -1, 1), (42, H - 42, 1, -1), (W - 42, H - 42, -1, -1)]:
        d.line([(cx, cy), (cx + dx * 30, cy)], fill=CYAN + (120,), width=3)
        d.line([(cx, cy), (cx, cy + dy * 30)], fill=CYAN + (120,), width=3)
    if plus:
        for px, py in [(W // 2, 150), (W // 2, H - 150), (300, H // 2), (W - 300, H // 2)]:
            d.line([(px - 9, py), (px + 9, py)], fill=(120, 160, 200, 40), width=2)
            d.line([(px, py - 9), (px, py + 9)], fill=(120, 160, 200, 40), width=2)

def rrect(d, box, r=20, fill=None, outline=None, width=2):
    d.rounded_rectangle(box, radius=r, fill=fill, outline=outline, width=width)

def chip(d, x, y, text, font, fg, bg=(8, 12, 18), border=None, pad=(18, 9)):
    tw = font.getlength(text)
    asc, desc = font.getmetrics()
    box = (x, y, int(x + tw + pad[0] * 2), y + asc + desc + pad[1] * 2 - 6)
    rrect(d, box, r=(box[3] - box[1]) // 2, fill=bg, outline=border, width=2 if border else 0)
    d.text((x + pad[0], y + pad[1] - 3), text, font=font, fill=fg)
    return box

def wrap(text, font, maxw):
    lines, line = [], ""
    for word in text.split():
        t = (line + " " + word).strip()
        if font.getlength(t) <= maxw: line = t
        else:
            if line: lines.append(line)
            line = word
    if line: lines.append(line)
    return lines

def text_block(d, xy, text, font, fill, maxw=None, lh=None, align="left", max_lines=None):
    x, y = xy
    asc, desc = font.getmetrics()
    lh = lh or int((asc + desc) * 1.32)
    lines = wrap(text, font, maxw) if maxw else [text]
    if max_lines and len(lines) > max_lines:
        print(f"  [!] overflow {len(lines)}>{max_lines}: {text[:56]}")
    for ln in lines:
        lx = x if align == "left" else x + (maxw - font.getlength(ln)) // 2
        d.text((lx, y), ln, font=font, fill=fill)
        y += lh
    return y - xy[1], lines

def halo(im, d, xy, text, font, rad=44, alpha=165):
    """Мягкая тёмная подложка за текстом — читаемость на любом фоне."""
    x, y = xy
    lay = Image.new("RGBA", im.size, (0, 0, 0, 0))
    ImageDraw.Draw(lay).text((x, y), text, font=font, fill=(4, 7, 11, alpha))
    im.paste(Image.alpha_composite(im.convert("RGBA"), lay.filter(ImageFilter.GaussianBlur(rad))).convert("RGB"), (0, 0))

def glow_text(im, d, xy, text, font, fill, glow=ORANGE, rad=26, alpha=110):
    x, y = xy
    lay = Image.new("RGBA", im.size, (0, 0, 0, 0))
    ImageDraw.Draw(lay).text((x, y), text, font=font, fill=glow + (alpha,))
    im.paste(Image.alpha_composite(im.convert("RGBA"), lay.filter(ImageFilter.GaussianBlur(rad))).convert("RGB"), (0, 0))
    d.text((x, y), text, font=font, fill=fill)

def header(d, kicker, title_lines, tsize=104, ky=96):
    d.text((100, ky), kicker, font=F("mono600", 26), fill=ORANGE)
    y = ky + 66
    for ln, col in title_lines:
        text_block(d, (100, y), ln, F("disp800", tsize), col)
        y += int(tsize * 1.28)
    return y

def footer(d, text_left, page=None):
    d.line([(100, H - 74), (W - 100, H - 74)], fill=LINE + (160,), width=2)
    d.text((100, H - 52), text_left, font=F("mono400", 21), fill=FAINT)
    if page:
        d.text((W - 100 - F("mono400", 21).getlength(page), H - 52), page, font=F("mono400", 21), fill=FAINT)

def bullet_list(d, x, y, items, font, maxw, gap=15, markc=ORANGE, textc=(205, 216, 228)):
    for it in items:
        d.text((x, y + int(font.size * 0.28)), "▸", font=F("body700", int(font.size * 0.86)), fill=markc)
        hgt, _ = text_block(d, (x + int(font.size * 1.15), y), it, font, textc, maxw=maxw - int(font.size * 1.15), lh=int(font.size * 1.3))
        y += max(hgt, font.size) + gap
    return y

# =====================================================================
# СЛАЙД 1 — ТИТУЛ (картинка видна справа полностью)
# =====================================================================
def slide1():
    im, d = canvas()
    cover_img(im, "hero_cover.png", (0, 0, W, H), brightness=1.30)
    hfade(im, (0, 0, 1000, H), BG, 246, 196)               # плотно слева,
    hfade(im, (1000, 0, 1750, H), BG, 196, 0)              # дальше растворяемся
    vgrad(im, (0, H - 460, W, H), BG, BG, alpha=175, flip=True)  # лёгкий низ под статы
    deco(d)

    d.rectangle((100, 88, 148, 136), fill=ORANGE)
    d.text((124, 112), "U", font=F("disp900", 34), fill=(10, 14, 20), anchor="mm")
    d.text((170, 96), "UNIT GROUP", font=F("disp700", 40), fill=WHITE)
    d.text((W - 100 - F("mono500", 24).getlength("МОСКВА И МО · 2026"), 104),
           "МОСКВА И МО · 2026", font=F("mono500", 24), fill=(196, 208, 222))

    chip(d, 100, 300, "ПРОЕКТИРОВАНИЕ · СТРОИТЕЛЬСТВО · ИНЖЕНЕРНЫЕ СИСТЕМЫ · ОТДЕЛКА",
         F("mono500", 24), CYAN, border=CYAN + (90,))

    for xy, t in [((96, 396), "ПОЛНЫЙ ЦИКЛ."), ((96, 556), "ОДНА ОТВЕТСТВЕННОСТЬ.")]:
        halo(im, d, xy, t, F("disp800", 116))
    glow_text(im, d, (96, 396), "ПОЛНЫЙ ЦИКЛ.", F("disp800", 116), WHITE, rad=30, alpha=55)
    glow_text(im, d, (96, 556), "ОДНА ОТВЕТСТВЕННОСТЬ.", F("disp800", 116), ORANGE2, rad=34, alpha=85)
    text_block(d, (100, 752), "От фундамента до космоса: группа из трёх компаний закрывает весь цикл работ — от буронабивной сваи до финишной отделки под ключ.",
               F("body500", 32), (208, 220, 233), maxw=1460, lh=46)

    py = H - 320
    overlay(im, (100, py, W - 100, py + 190), (12, 20, 33), 205)
    d.line([(100, py), (W - 100, py)], fill=ORANGE + (200,), width=5)
    stats = [("3", "компании — единый контур"),
             ("10+", "лет спецработ: метро, тоннели"),
             ("288 м", "НКЦ — флагманский объект"),
             ("24/7", "работы в любом климате")]
    cw = (W - 200) / 4
    for i, (num, lab) in enumerate(stats):
        x = 100 + i * cw
        if i: d.line([(int(x), py + 34), (int(x), py + 156)], fill=LINE + (220,), width=2)
        d.text((int(x + 44), py + 30), num, font=F("mono700", 62), fill=ORANGE if i in (0, 2) else WHITE)
        text_block(d, (int(x + 44), py + 112), lab, F("body500", 24), (178, 190, 205), maxw=cw - 70, lh=31, max_lines=2)

    footer(d, "ООО «ЮНИТ ИНЖИНИРИНГ»  ·  ООО «АММ-СТРОЙ»  ·  ООО «МОСТЫ И ТОННЕЛИ»", page="01 / 05")
    return im

# =====================================================================
# СЛАЙД 2 — СТРУКТУРА (появился фоновый панорамный кадр, сетки нет)
# =====================================================================
def slide2():
    im, d = canvas()
    cover_img(im, "bg_city2.png", (0, 0, W, H), brightness=1.05, extra_dark=86)
    vgrad(im, (0, 0, W, 430), BG, BG, alpha=178)
    vgrad(im, (0, H - 420, W, H), BG, BG, alpha=170, flip=True)
    deco(d)

    header(d, "01 · СТРУКТУРА ГРУППЫ", [("Три компании. Один контур.", WHITE)])
    tf = F("disp800", 104)
    suby = 162 + sum(tf.getmetrics()) + 22          # ниже заголовка с запасом
    text_block(d, (100, suby), "Каждая компания закрывает свой сегмент, заказчик получает один договор и одну ответственность за весь цикл — от проекта до отделки.",
               F("body500", 30), (196, 208, 222), maxw=1900, lh=42)

    cards = [
        dict(role="ГОЛОВНАЯ · ИНЖИНИРИНГ", rolec=ORANGE, name="ЮНИТ\nИНЖИНИРИНГ",
             items=["Проектирование и строительство любой сложности",
                    "Комплекс работ на объекте НКЦ — башня 288 м, 257 000 м²",
                    "Капремонт тоннеля «Новокутузовский» на ТТК",
                    "Стрелковый клуб «Румянцево» — проект и стройка"],
             head="Махадов Ильян Абдурагимович", meta="гендиректор · с 2025 г."),
        dict(role="ОПЕРАЦИОННАЯ · СТРОЙКА", rolec=CYAN, name="АММ-СТРОЙ",
             items=["Отделка «под ключ»: плитка, паркет, ГКЛ, лепнина",
                    "Демонтаж и алмазная резка ж/б любой сложности",
                    "Черновые работы, столярка, двери, плинтусы",
                    "Сантехника, электрика, ОВиК и газоснабжение"],
             head="Михеев Михаил Вахитович", meta="гендиректор · с 05.2026 · команда AIVA"),
        dict(role="СПЕЦПОДРЯД · 10+ ЛЕТ", rolec=GOLD, name="МОСТЫ\nИ ТОННЕЛИ",
             items=["Мосты, тоннели, метрополитен, ИССО",
                    "БНС · БСС · БКС · «стена в грунте» Ø620–1200 мм",
                    "Монолит и гидроизоляция полного цикла",
                    "Собственный парк: Casagrande · Hilti · работа 24/7"],
             head="Хидешели Леван Иванович", meta="гендиректор · ooomit.com"),
    ]
    top, ch_ = 400, 668
    cw = int((W - 200 - 80) // 3)
    for i, c in enumerate(cards):
        x = 100 + i * (cw + 40)
        box = (x, top, x + cw, top + ch_)
        overlay(im, box, CARD2, 235)
        rrect(d, box, r=24, outline=LINE + (255,), width=3)
        d.line([(x + 24, top), (x + 190, top)], fill=c["rolec"] + (240,), width=6)
        chip(d, x + 28, top + 34, c["role"], F("mono600", 21), c["rolec"], border=c["rolec"] + (110,))
        ny = top + 106
        for ln in c["name"].split("\n"):
            text_block(d, (x + 28, ny), ln, F("disp700", 44), WHITE, max_lines=1)
            ny += 58
        ny += 10
        d.line([(x + 28, ny), (x + cw - 28, ny)], fill=LINE + (200,), width=2)
        ny += 24
        bullet_list(d, x + 28, ny, c["items"], F("body500", 25), cw - 56)
        by = top + ch_ - 118
        overlay(im, (x + 1, by, x + cw - 1, top + ch_ - 1), (12, 19, 30), 215)
        text_block(d, (x + 28, by + 18), c["head"], F("body600", 25), WHITE, maxw=cw - 56, max_lines=1)
        text_block(d, (x + 28, by + 58), c["meta"], F("mono400", 20), GREY, maxw=cw - 56, max_lines=1)

    ly = top + ch_ + 42
    overlay(im, (100, ly, W - 100, ly + 90), CARD, 220)
    d.line([(100, ly), (W - 100, ly)], fill=ORANGE + (190,), width=4)
    t = "ОДИН ТЕНДЕР  →  ОДИН ДОГОВОР  →  ОДНА ОТВЕТСТВЕННОСТЬ"
    f = F("mono700", 34)
    d.text(((W - f.getlength(t)) / 2, ly + 25), t, font=f, fill=WHITE)
    footer(d, "UNIT GROUP · структура группы", page="02 / 05")
    return im

# =====================================================================
# СЛАЙД 3 — КОМПЕТЕНЦИИ (буровая видна целиком, есть фон)
# =====================================================================
def slide3():
    im, d = canvas()
    cover_img(im, "bg_city2.png", (0, 0, W, H), brightness=1.0, extra_dark=96)
    vgrad(im, (0, 0, W, 420), BG, BG, alpha=185)
    vgrad(im, (0, H - 300, W, H), BG, BG, alpha=180, flip=True)
    deco(d)

    header(d, "02 · КОМПЕТЕНЦИИ", [("Полный цикл. Без стыков.", WHITE)])
    tf = F("disp800", 104)
    suby = 162 + sum(tf.getmetrics()) + 22
    text_block(d, (100, suby), "Шесть этапов, которые на рынке обычно делят между 3–5 подрядчиками. У нас их закрывает одна команда, одна смета и одна ответственность.",
               F("body500", 30), (196, 208, 222), maxw=1900, lh=42)

    py = 400
    stages = [("01", "Проектирование", "стадии П и Р, BIM, сметы"),
              ("02", "Буровые работы", "БНС, БСС, БКС, «стена в грунте»"),
              ("03", "Монолит", "армирование, опалубка, бетон"),
              ("04", "Гидроизоляция", "2–4 слоя, металл, инъекции"),
              ("05", "Инженерные системы", "ОВиК, вода, газ, электрика"),
              ("06", "Отделка под ключ", "от стяжки до дверей и света")]
    sw = int((W - 200 - 5 * 26) // 6)
    for i, (n, t, s) in enumerate(stages):
        x = 100 + i * (sw + 26)
        box = (x, py, x + sw, py + 200)
        overlay(im, box, CARD2, 240)
        rrect(d, box, r=18, outline=LINE + (255,), width=2)
        d.line([(x, py), (x + 62, py)], fill=ORANGE + (230,), width=5)
        d.text((x + 22, py + 20), n, font=F("mono700", 34), fill=ORANGE)
        text_block(d, (x + 22, py + 76), t, F("body700", 27), WHITE, maxw=sw - 44, lh=33, max_lines=2)
        text_block(d, (x + 22, py + 146), s, F("body500", 21), (178, 190, 205), maxw=sw - 44, lh=26, max_lines=2)
        if i < 5:
            d.polygon([(x + sw + 5, py + 100), (x + sw + 19, py + 89), (x + sw + 19, py + 111)], fill=ORANGE + (220,))

    ty, th_ = 660, 480
    lb = (100, ty, 1060, ty + th_)
    fit_contain(im, "drill_rig.png", lb, brightness=1.24)
    rrect(d, lb, r=22, outline=LINE + (255,), width=3)
    chip(d, lb[0] + 26, lb[1] + 24, "СОБСТВЕННЫЙ ПАРК ТЕХНИКИ", F("mono600", 21), WHITE, border=ORANGE + (160,))
    vgrad(im, (lb[0], lb[3] - 250, lb[2], lb[3]), BG, BG, alpha=228, flip=True)
    text_block(d, (lb[0] + 26, lb[3] - 196), "Буровые установки Casagrande и канатная пила Hilti",
               F("disp700", 34), WHITE, maxw=lb[2] - lb[0] - 52, lh=44, max_lines=2)
    text_block(d, (lb[0] + 26, lb[3] - 102), "Свайные работы Ø620–1200 мм в грунтах до 10 группы, в напорных водоносных горизонтах. Мобильные участки — круглосуточно.",
               F("body500", 23), (200, 212, 226), maxw=lb[2] - lb[0] - 52, lh=30, max_lines=3)

    sx = 1100
    sw2 = W - 100 - sx
    sh = int((th_ - 2 * 24) // 3)
    rows = [("291 БКС", "ограждение котлована ЖК — Шмитовский проезд, ЦАО"),
            ("1 000 т", "статические испытания свай — ТПУ «Мичуринский проспект»"),
            ("20 000 м³", "монолит основных конструкций — ст. «Кленовый бульвар», БКЛ")]
    for i, (num, lab) in enumerate(rows):
        y = ty + i * (sh + 24)
        box = (sx, y, W - 100, y + sh)
        overlay(im, box, CARD2, 240)
        rrect(d, box, r=20, outline=LINE + (255,), width=2)
        d.line([(sx, y), (sx, y + 84)], fill=ORANGE + (220,), width=5)
        d.text((sx + 34, y + 18), num, font=F("mono700", 56), fill=ORANGE)
        text_block(d, (sx + 34, y + 96), lab, F("body500", 24), (205, 216, 230), maxw=sw2 - 68, lh=31, max_lines=2)

    footer(d, "UNIT GROUP · компетенции", page="03 / 05")
    return im

# =====================================================================
# СЛАЙД 4 — ПРОЕКТЫ (все изображения целиком + фото в каждой карточке)
# =====================================================================
def slide4():
    im, d = canvas()
    cover_img(im, "bg_city2.png", (0, 0, W, H), brightness=1.0, extra_dark=100)
    vgrad(im, (0, 0, W, 380), BG, BG, alpha=188)
    vgrad(im, (0, H - 320, W, H), BG, BG, alpha=185, flip=True)
    deco(d)

    header(d, "03 · ПРОЕКТЫ-ДОКАЗАТЕЛЬСТВА", [("Там, где нас уже проверили.", WHITE)], tsize=96)

    top = 330
    # флагман НКЦ: фото сверху БЕЗ кропа (зона ровно 16:9), текст снизу
    fb = (100, top, 940, top + 820)
    img_h = int((fb[2] - fb[0]) / (1672 / 941))          # ~472
    cover_img(im, "nkz_tower.png", (fb[0], fb[1], fb[2], fb[1] + img_h), brightness=1.22)
    rrect(d, (fb[0], fb[1], fb[2], fb[1] + img_h), r=22, outline=LINE + (255,), width=3)
    chip(d, fb[0] + 26, fb[1] + 24, "ФЛАГМАН · КОСМОС", F("mono600", 22), WHITE, border=ORANGE + (170,))
    # текстовая зона карточки
    tb = (fb[0], fb[1] + img_h, fb[2], fb[3])
    overlay(im, (tb[0] + 1, tb[1], tb[2] - 1, tb[3]), CARD2, 245)
    d.line([(tb[0] + 1, tb[1]), (tb[2] - 1, tb[1])], fill=ORANGE + (170,), width=3)
    text_block(d, (tb[0] + 28, tb[1] + 22), "Национальный космический центр",
               F("disp700", 40), WHITE, maxw=tb[2] - tb[0] - 56, lh=50, max_lines=2)
    text_block(d, (tb[0] + 28, tb[1] + 126), "Комплекс работ на объекте совместного проекта Правительства Москвы и «Роскосмоса». Открыт 13.09.2025.",
               F("body500", 23), (198, 210, 224), maxw=tb[2] - tb[0] - 56, lh=30, max_lines=3)
    sy = tb[3] - 118
    d.line([(tb[0] + 28, sy - 16), (tb[2] - 28, sy - 16)], fill=LINE + (220,), width=2)
    cols = [("288 м", "47 этажей"), ("257 000", "м² площади"), ("20 000", "раб. мест")]
    cwid = (tb[2] - tb[0] - 56) / 3
    for i, (n, l) in enumerate(cols):
        x = tb[0] + 28 + i * cwid
        d.text((int(x), sy), n, font=F("mono700", 42), fill=ORANGE)
        text_block(d, (int(x), sy + 58), l, F("body500", 22), GREY, maxw=cwid - 12, max_lines=1)

    # правая сетка 2x2 — теперь ВО ВСЕХ карточках есть фото (целиком)
    gx0 = 980
    gw = int((W - 100 - gx0 - 30) // 2)
    gh = int((820 - 30) // 2)
    cells = [
        dict(img="metro_build.png", tag="МЕТРО · БКЛ И ТРОИЦКАЯ", title="4 станции метро",
             text="Говорово · Столбово · «Кленовый бульвар» (20 000 м³) · вестибюль «Академической», 2025"),
        dict(img="tunnel_road.png", tag="ТОННЕЛИ И ДОРОГИ", title="Тоннель в Видном — 300 м",
             text="Секция трассы «ЮЛА» (концессия 86,7 млрд ₽) · капремонт Новокутузовского тоннеля, ТТК"),
        dict(img="piles_geo.png", tag="ГЕОТЕХНИКА · СВАИ", title="291 БКС и 238 БНС",
             text="Ограждение котлована — Шмитовский проезд, ЦАО · свайное основание УДС ул. Маршала Шестопалова"),
        dict(img="towers_dev.png", tag="ДЕВЕЛОПМЕНТ · ТПУ", title="«Мичуринский», Башня А",
             text="Свайное основание БНС Ø1000–1200 мм и испытания нагрузкой 1 000 т в составе ТПУ"),
    ]
    for i, c in enumerate(cells):
        cx = gx0 + (i % 2) * (gw + 30)
        cy = top + (i // 2) * (gh + 30)
        box = (cx, cy, cx + gw, cy + gh)
        fit_contain(im, c["img"], box, brightness=1.18)
        rrect(d, box, r=20, outline=LINE + (255,), width=3)
        vgrad(im, (box[0], box[3] - 205, box[2], box[3]), BG, BG, alpha=238, flip=True)
        chip(d, box[0] + 22, box[1] + 20, c["tag"], F("mono600", 19), CYAN)
        text_block(d, (box[0] + 22, box[3] - 166), c["title"], F("body700", 30), WHITE, maxw=gw - 44, lh=38, max_lines=2)
        text_block(d, (box[0] + 22, box[3] - 118), c["text"], F("body500", 21), (186, 198, 213), maxw=gw - 44, lh=27, max_lines=3)

    footer(d, "Группа выполняла комплекс работ на объектах; генпроектировщик и генподрядчик ряда проектов — иные организации.", page="04 / 05")
    return im

# =====================================================================
# СЛАЙД 5 — КОНТАКТЫ
# =====================================================================
def slide5():
    im, d = canvas()
    cover_img(im, "hero_cover.png", (0, 0, W, H), brightness=1.30)
    overlay(im, (0, 0, W, H), BG, 178)
    vgrad(im, (0, 0, W, 460), BG, BG, alpha=150)
    deco(d)

    header(d, "04 · КОНТАКТЫ", [("Готовы к диалогу.", WHITE), ("Обсудим ваш объект.", ORANGE2)], tsize=92)
    text_block(d, (100, 418), "Полный цикл. Одна ответственность. От котлована и свай до отделки под ключ — в одном договоре.",
               F("body500", 30), (205, 216, 230), maxw=1560, lh=42)

    top, ch_ = 528, 536
    cw = int((W - 200 - 80) // 3)
    cards = [
        dict(role="ГОЛОВНАЯ · EPC", rolec=ORANGE, name="ООО «ЮНИТ ИНЖИНИРИНГ»",
             rows=[("Гендиректор", "Махадов Ильян Абдурагимович"),
                   ("Телефон", "+7 915 186-66-11 · +7 916 839-40-20"),
                   ("E-mail", "ooo-unit@list.ru"),
                   ("Адрес", "105082, Москва, ул. Б. Почтовая, 26В, стр. 1"),
                   ("ИНН / ОГРН", "9701308826 / 1257700109072")]),
        dict(role="ОПЕРАЦИОННЫЙ КОНТУР", rolec=CYAN, name="ООО «АММ-СТРОЙ»",
             rows=[("Гендиректор", "Михеев Михаил Вахитович"),
                   ("E-mail", "mikheev.mikhail13@mail.ru"),
                   ("Адрес", "117393, Москва, ул. Архитектора Власова, 55"),
                   ("Создана", "04.05.2026 · команда и портфель AIVA"),
                   ("ИНН / ОГРН", "9729419389 / 1267700156052")]),
        dict(role="СПЕЦПОДРЯД · 10+ ЛЕТ", rolec=GOLD, name="ООО «МОСТЫ И ТОННЕЛИ»",
             rows=[("Гендиректор", "Хидешели Леван Иванович"),
                   ("Телефон", "+7 499 148-63-83"),
                   ("Сайт / e-mail", "ooomit.com · info@ooomit.com"),
                   ("Адрес", "МО, Люберцы, ул. Электрификации, 3, стр. 3"),
                   ("ИНН / ОГРН", "5001107429 / 1165001050785")]),
    ]
    for i, c in enumerate(cards):
        x = 100 + i * (cw + 40)
        box = (x, top, x + cw, top + ch_)
        overlay(im, box, CARD2, 240)
        rrect(d, box, r=24, outline=LINE + (255,), width=3)
        d.line([(x + 24, top), (x + 190, top)], fill=c["rolec"] + (240,), width=6)
        chip(d, x + 28, top + 32, c["role"], F("mono600", 21), c["rolec"], border=c["rolec"] + (110,))
        text_block(d, (x + 28, top + 94), c["name"], F("disp700", 36), WHITE, maxw=cw - 56, lh=46, max_lines=2)
        y = top + 170
        for k, v in c["rows"]:
            d.text((x + 28, y), k.upper(), font=F("mono500", 20), fill=GREY2)
            hgt, _ = text_block(d, (x + 28, y + 28), v, F("body600", 26), (218, 228, 238), maxw=cw - 56, lh=33, max_lines=2)
            y += 28 + 33 * max(1, len(wrap(v, F("body600", 26), cw - 56))) + 6

    footer(d, "«Ваш надёжный партнёр в мире инженерии»  ·  UNIT GROUP, 2026", page="05 / 05")
    return im

# =====================================================================
def main():
    out = os.path.join(BASE, "slides")
    os.makedirs(out, exist_ok=True)
    pages = [slide1(), slide2(), slide3(), slide4(), slide5()]
    paths = []
    for i, p in enumerate(pages, 1):
        fp = os.path.join(out, f"slide{i}.png")
        p.save(fp, optimize=True)
        paths.append(fp)
        print("saved", fp)
    pdf_pages = [p.convert("RGB") for p in pages]
    pdf = os.path.join(BASE, "UNIT_GROUP_Презентация.pdf")
    pdf_pages[0].save(pdf, save_all=True, append_images=pdf_pages[1:], resolution=180.0, quality=93)
    print("PDF:", pdf, os.path.getsize(pdf) // 1024, "KB")
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]
    for fp in paths:
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(fp, 0, 0, width=Emu(12192000), height=Emu(6858000))
    prs.core_properties.title = "UNIT GROUP — инженерно-строительная группа полного цикла"
    prs.core_properties.author = "UNIT GROUP"
    pptx = os.path.join(BASE, "UNIT_GROUP_Презентация.pptx")
    prs.save(pptx)
    add_transitions(pptx)
    print("PPTX:", pptx, os.path.getsize(pptx) // 1024, "KB")


def add_transitions(pptx_path):
    """Переходы между слайдами: 1-й — fade, остальные — Morph (плавная
    трансформация) с fade-фолбэком для старых PowerPoint. Инъекция XML —
    единственный рабочий способ, т.к. у python-pptx нет API переходов."""
    from pptx import Presentation
    from lxml import etree
    P  = "http://schemas.openxmlformats.org/presentationml/2006/main"
    MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    P14  = "http://schemas.microsoft.com/office/powerpoint/2010/main"
    P159 = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"
    fade = (f'<p:transition xmlns:p="{P}" spd="slow">'
            f'<p:fade/></p:transition>')
    morph = (f'<mc:AlternateContent xmlns:mc="{MC}">'
             f'<mc:Choice xmlns:p159="{P159}" Requires="p159">'
             f'<p:transition xmlns:p="{P}" xmlns:p14="{P14}" spd="slow" p14:dur="1300">'
             f'<p159:morph option="byObject"/></p:transition></mc:Choice>'
             f'<mc:Fallback>'
             f'<p:transition xmlns:p="{P}" spd="slow"><p:fade/></p:transition>'
             f'</mc:Fallback></mc:AlternateContent>')
    prs = Presentation(pptx_path)
    for i, slide in enumerate(prs.slides):
        slide.element.append(etree.fromstring(fade if i == 0 else morph))
    prs.save(pptx_path)
    print("transitions: 1x fade + 4x morph(fallback fade)")

if __name__ == "__main__":
    main()
