# -*- coding: utf-8 -*-
"""
UNIT GROUP — ЭКСПЕРИМЕНТАЛЬНАЯ АНИМИРОВАННАЯ версия PPTX («ва-банк»).

Идея: слайды больше не «одна картинка». Каждый слайд нарезается на слои
(статичный фон + анимируемые объекты) методом диф-рендера: слайд-функция
из build_unit_group.py прогоняется несколько раз с лимитом графических
примитивов, разница между последовательными прогонами = слой с альфой.
В PowerPoint слои складываются в исходный кадр, но появляются каскадом
эффектами Float In / Fade (инъекция p:timing XML — python-pptx API нет).
Morph-переходы между слайдами сохранены. Совпадение с PDF проверяется
численно (композит слоёв == полный рендер).

Риск: если PowerPoint закапризничает — у нас есть безопасный файл
UNIT GROUP_Презентация.pptx без покадровых анимаций.
"""
import os, sys, copy
from PIL import Image, ImageDraw, ImageChops, ImageFilter

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE)
import build_unit_group as B
from build_unit_group import W, H, F  # noqa

STATE = {"n": 0, "limit": 10 ** 12, "log": []}

class CImg:
    """Прокси над Image: считает paste-операции, умеет «выключать» после лимита."""
    def __init__(self, im):
        object.__setattr__(self, "_im", im)
    def paste(self, *a, **k):
        STATE["n"] += 1
        STATE["log"].append((STATE["n"], "paste", str(a[1]) if len(a) > 1 else ""))
        if STATE["n"] <= STATE["limit"]:
            object.__getattribute__(self, "_im").paste(*a, **k)
    def __getattr__(self, name):
        return getattr(object.__getattribute__(self, "_im"), name)

class CDraw:
    """Прокси над ImageDraw: считает все примитивы, после лимита — no-op."""
    def __init__(self, d):
        object.__setattr__(self, "_d", d)
    def __getattr__(self, name):
        f = getattr(object.__getattribute__(self, "_d"), name)
        if not callable(f):
            return f
        def wrapper(*a, **k):
            STATE["n"] += 1
            info = ""
            if name == "text" and len(a) >= 2 and isinstance(a[1], str):
                info = a[1][:44]
            elif a:
                info = str(a[0])[:44]
            STATE["log"].append((STATE["n"], name, info))
            if STATE["n"] <= STATE["limit"]:
                return f(*a, **k)
        return wrapper

_orig_canvas = B.canvas
def patched_canvas():
    im = Image.new("RGB", (W, H), B.BG)
    return CImg(im), CDraw(ImageDraw.Draw(im))
B.canvas = patched_canvas

def render_with_limit(fn, limit, want_log=False):
    STATE["n"] = 0
    STATE["limit"] = limit
    STATE["log"] = []
    result = fn()
    im = result[0] if isinstance(result, tuple) else result
    real = object.__getattribute__(im, "_im")
    return (real, list(STATE["log"])) if want_log else real

def diff_layer(before, after, dilate=3):
    """Слой = пиксели, изменившиеся между прогонами (альфа-маска)."""
    d = ImageChops.difference(before.convert("L"), after.convert("L"))
    mask = d.point(lambda v: 255 if v > 5 else 0)
    if dilate:
        mask = mask.filter(ImageFilter.MaxFilter(dilate))
    layer = after.convert("RGBA")
    layer.putalpha(mask)
    return layer

def composite(base, layers):
    acc = base.convert("RGBA")
    for l in layers:
        acc = Image.alpha_composite(acc, l)
    return acc.convert("RGB")

# ---------------------------------------------------------------- timings
NS = ("xmlns:p=\"http://schemas.openxmlformats.org/presentationml/2006/main\" "
      "xmlns:r=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships\"")

def _effect_par(ctr, spid, delay, dur, effect, node_type):
    """Один эффект-контейнер: set visible + fade (+ ppt_y для float)."""
    i1, i2, i3, i4 = ctr, ctr + 1, ctr + 2, ctr + 3
    extra = ""
    if effect == "float":
        extra = f'''<p:anim calcmode="lin" valueType="num">
          <p:cBhvr additive="base">
            <p:cTn id="{i4}" dur="{dur}" fill="hold"/>
            <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
            <p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst>
          </p:cBhvr>
          <p:tavLst>
            <p:tav tm="0"><p:val><p:strVal val="#ppt_y+0.028"/></p:val></p:tav>
            <p:tav tm="100000"><p:val><p:strVal val="#ppt_y"/></p:val></p:tav>
          </p:tavLst>
        </p:anim>'''
    return f'''<p:par>
        <p:cTn id="{i1}" presetID="10" presetClass="entr" presetSubtype="0"
               fill="hold" grpId="0" nodeType="{node_type}">
          <p:stCondLst><p:cond delay="{delay}"/></p:stCondLst>
          <p:childTnLst>
            <p:set>
              <p:cBhvr>
                <p:cTn id="{i2}" dur="1" fill="hold">
                  <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                </p:cTn>
                <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
                <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
              </p:cBhvr>
              <p:to><p:strVal val="visible"/></p:to>
            </p:set>
            <p:animEffect transition="in" filter="fade">
              <p:cBhvr>
                <p:cTn id="{i3}" dur="{dur}"/>
                <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
              </p:cBhvr>
            </p:animEffect>
            {extra}
          </p:childTnLst>
        </p:cTn>
      </p:par>'''

def timing_xml(effects):
    """mainSeq с автостартом (первый эффект afterPrevious, остальные withPrevious)."""
    ctr = [10]
    pars = []
    for k, (spid, delay, dur, effect) in enumerate(effects):
        nt = "afterEffect" if k == 0 else "withEffect"
        pars.append(_effect_par(ctr[0], spid, delay, dur, effect, nt))
        ctr[0] += 10
    blds = "".join(
        f'<p:bldP spid="{spid}" grpId="0"/>' for spid, *_ in effects)
    return f'''<p:timing {NS}>
      <p:tnLst>
        <p:par>
          <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
            <p:childTnLst>
              <p:seq concurrent="1" nextAc="seek">
                <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
                  <p:childTnLst>
                    <p:par>
                      <p:cTn id="3" fill="hold">
                        <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                        <p:childTnLst>
                          <p:par>
                            <p:cTn id="4" fill="hold">
                              <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                              <p:childTnLst>
                                {''.join(pars)}
                              </p:childTnLst>
                            </p:cTn>
                          </p:par>
                        </p:childTnLst>
                      </p:cTn>
                    </p:par>
                  </p:childTnLst>
                </p:cTn>
                <p:prevCondLst>
                  <p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>
                </p:prevCondLst>
                <p:nextCondLst>
                  <p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>
                </p:nextCondLst>
              </p:seq>
            </p:childTnLst>
          </p:cTn>
        </p:par>
      </p:tnLst>
      <p:bldLst>{blds}</p:bldLst>
    </p:timing>'''

# ---------------------------------------------------------------- план слоёв
# На каждый слайд: список (имя, маркер_начала, эффект, dur, задержка_от_предыдущего)
# маркер = (kind, подстрока) — первый примитив, чей лог содержит подстроку.
PLAN = {
    1: [
        ("slogan",   ("text", "ПОЛНЫЙ ЦИКЛ."),        "fade",  1100, 200),
        ("subtext",  ("text", "От фундамента"),       "fade",   900, 150),
        ("stats",    ("paste", ""),                    "float", 900, 200),   # берём по индексу ниже
    ],
    2: [
        ("card1", ("text", "ГОЛОВНАЯ"),        "float", 800, 150),
        ("card2", ("text", "ОПЕРАЦИОННАЯ"),    "float", 800, 300),
        ("card3", ("text", "СПЕЦПОДРЯД"),      "float", 800, 300),
        ("band",  ("text", "ОДИН ТЕНДЕР"),     "fade",  700, 350),
    ],
    3: [
        ("stages", ("text", "Проектирование"), "fade",  800, 150),
        ("drill",  ("paste", ""),              "fade",  900, 250),
        ("stats",  ("text", "291 БКС"),        "float", 900, 250),
    ],
    4: [
        ("nkzc",  ("text", "ФЛАГМАН"),          "float", 900, 200),
        ("cell1", ("text", "МЕТРО"),            "fade",  800, 250),
        ("cell2", ("text", "ТОННЕЛИ И ДОРОГИ"), "fade",  800, 200),
        ("cell3", ("text", "ГЕОТЕХНИКА"),       "fade",  800, 200),
        ("cell4", ("text", "ДЕВЕЛОПМЕНТ"),      "fade",  800, 200),
    ],
    5: [
        ("card1", ("text", "ГОЛОВНАЯ"),            "float", 800, 200),
        ("card2", ("text", "ОПЕРАЦИОННЫЙ КОНТУР"), "float", 800, 300),
        ("card3", ("text", "СПЕЦПОДРЯД"),          "float", 800, 300),
    ],
}
# ручные индексные маркеры там, где текст не уникально найти (slide1 stats, slide3 drill)
INDEX_HINTS = {1: {"stats": "after_text:24/7"}, 3: {"drill": "paste_before:291 БКС"}}

SLIDE_FNS = {1: B.slide1, 2: B.slide2, 3: B.slide3, 4: B.slide4, 5: B.slide5}

def find_marker_idx(log, kind, needle, start=0):
    for idx, k, info in log:
        if idx <= start:
            continue
        if k == kind and needle in info:
            return idx
    return None

def build_layers():
    """Возвращает {slide: (base_png_path, [layer_paths])}."""
    outdir = os.path.join(BASE, "anim_layers")
    os.makedirs(outdir, exist_ok=True)
    result = {}
    for sn, plan in PLAN.items():
        fn = SLIDE_FNS[sn]
        full, log = render_with_limit(fn, 10 ** 12, want_log=True)
        # индексы начала слоёв
        boundaries = []
        start = 0
        for name, (kind, needle), effect, dur, delay in plan:
            if name in INDEX_HINTS.get(sn, {}):
                hint = INDEX_HINTS[sn][name]
                if hint.startswith("after_text:"):
                    anchor = find_marker_idx(log, "text", hint.split(":", 1)[1])
                    boundaries.append((name, anchor, effect, dur, delay))
                elif hint.startswith("paste_before:"):
                    anchor = find_marker_idx(log, "text", hint.split(":", 1)[1])
                    # слой начинается с paste, идущего ПЕРЕД этим текстом
                    prev_paste = max([i for i, k, _ in log if k == "paste" and i < anchor], default=0)
                    boundaries.append((name, prev_paste + 1, effect, dur, delay))
                continue
            idx = find_marker_idx(log, kind, needle, start=start)
            if idx is None:
                raise RuntimeError(f"slide{sn} layer '{name}': marker {kind}:{needle} not found")
            boundaries.append((name, idx, effect, dur, delay))
            start = idx
        # base = всё до первого слоя; последний слой до конца
        cuts = [1] + [b[1] for b in boundaries] + [10 ** 12]
        base = render_with_limit(fn, cuts[1] - 1)
        base_path = os.path.join(outdir, f"s{sn}_base.png")
        base.save(base_path)
        layer_paths = []
        for li, (name, idx, effect, dur, delay) in enumerate(boundaries):
            before = render_with_limit(fn, idx - 1)
            after = render_with_limit(fn, cuts[li + 2] - 1)
            layer = diff_layer(before, after)
            p = os.path.join(outdir, f"s{sn}_l{li+1}_{name}.png")
            layer.save(p)
            layer_paths.append((p, effect, dur, delay, name))
        # проверка: композит == полный рендер
        comp = composite(base, [Image.open(p) for p, *_ in layer_paths])
        diff = ImageChops.difference(comp, full).convert("L")
        mean = sum(diff.getdata()) / (W * H)
        nonzero = sum(1 for v in diff.getdata() if v > 12)
        print(f"slide{sn}: layers={len(layer_paths)} mean_diff={mean:.2f} pixels>12: {nonzero}")
        result[sn] = (base_path, layer_paths, full)
    return result

def build_pptx():
    from pptx import Presentation
    from pptx.util import Emu
    from lxml import etree
    layers = build_layers()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]
    for sn in sorted(layers):
        base_path, lps, _ = layers[sn]
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(base_path, 0, 0, width=Emu(12192000), height=Emu(6858000))
        spids = []
        for p, effect, dur, delay, name in lps:
            pic = s.shapes.add_picture(p, 0, 0, width=Emu(12192000), height=Emu(6858000))
            spids.append((pic.shape_id, effect, dur, delay, name))
        # тайминги: задержки накапливаются от предыдущего эффекта
        effects, t = [], 0
        for spid, effect, dur, delay, name in spids:
            t += delay
            effects.append((spid, t, dur, effect))
        # порядок по схеме: transition ДО timing
        s._element.append(etree.fromstring(_transition_xml(sn)))
        s._element.append(etree.fromstring(timing_xml(effects)))
        print(f"slide{sn}: animated shapes:", [(sp, nm) for sp, _, _, _, nm in spids])
    prs.core_properties.title = "UNIT GROUP — анимированная версия (эксперимент)"
    prs.core_properties.author = "UNIT GROUP"
    out = os.path.join(BASE, "UNIT_GROUP_АНИМИРОВАННАЯ_эксперимент.pptx")
    prs.save(out)
    print("saved", out, os.path.getsize(out) // 1024, "KB")
    return out


def _transition_xml(sn):
    P  = "http://schemas.openxmlformats.org/presentationml/2006/main"
    MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    P14  = "http://schemas.microsoft.com/office/powerpoint/2010/main"
    P159 = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"
    if sn == 1:
        return f'<p:transition xmlns:p="{P}" spd="slow"><p:fade/></p:transition>'
    return (f'<mc:AlternateContent xmlns:mc="{MC}">'
            f'<mc:Choice xmlns:p159="{P159}" Requires="p159">'
            f'<p:transition xmlns:p="{P}" xmlns:p14="{P14}" spd="slow" p14:dur="1300">'
            f'<p159:morph option="byObject"/></p:transition></mc:Choice>'
            f'<mc:Fallback><p:transition xmlns:p="{P}" spd="slow"><p:fade/></p:transition>'
            f'</mc:Fallback></mc:AlternateContent>')

if __name__ == "__main__":
    build_pptx()
